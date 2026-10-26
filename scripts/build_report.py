"""Build the Team 3 DOCX from saved simulation results.

Run after scripts/create_report_figures.py. PDF export is a separate Word or
LibreOffice step; inspect the rendered pages before distributing either file.
"""
from pathlib import Path
import csv
import json
import re
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / 'report/work'
OUT = ROOT / 'report/Team_03.docx'
REPO = 'https://github.com/itsmePrithviraj/ia3-data-analytics'
WORK.mkdir(parents=True, exist_ok=True)
s = json.loads((ROOT / 'outputs/simulated/statistics.json').read_text())
if s.get('data_origin') != 'SIMULATED':
    raise ValueError('This report narrative is for the simulation; review it before using real results.')
with (ROOT / 'outputs/simulated/descriptive_statistics.csv').open() as f:
    summary = {r['variable']: r for r in csv.DictReader(f)}
with (ROOT / 'outputs/simulated/cleaned.csv').open() as f:
    rows = list(csv.DictReader(f))
n = s['n']
counts = {k: sum(r['perception_group'] == k for r in rows) for k in
          ['Underestimated >30 min', 'Within 30 min', 'Overestimated >30 min']}
f2 = lambda x: f'{float(x):.2f}'

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.2677); sec.page_height = Inches(11.6929)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.5)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    st = doc.styles[name]
    st.font.name = 'Times New Roman'; st.font.color.rgb = RGBColor(0, 0, 0)
    fonts = st.element.get_or_add_rPr().find(qn('w:rFonts'))
    for attribute in list(fonts.attrib):
        if 'theme' in attribute.lower(): del fonts.attrib[attribute]
    for script in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:'+script), 'Times New Roman')
    for border in st.element.findall('.//' + qn('w:pBdr')):
        border.getparent().remove(border)
    st.paragraph_format.widow_control = True
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for name, size in [('Heading 1',14), ('Heading 2',12)]:
    st = doc.styles[name]; st.font.size = Pt(size); st.font.bold = True
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.space_before = Pt(10); st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.keep_with_next = True
st = doc.styles['Title']; st.font.size = Pt(18); st.font.bold = True
st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
st.paragraph_format.line_spacing = 1.05
st.paragraph_format.space_before = Pt(7); st.paragraph_format.space_after = Pt(9)
cap = doc.styles['Caption']; cap.font.size = Pt(9); cap.font.italic = False
cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
cap.paragraph_format.line_spacing = 1.05
cap.paragraph_format.space_before = Pt(4); cap.paragraph_format.space_after = Pt(8)
header = sec.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.add_run('CES 2103  |  DATA ANALYTICS IA3  |  TEAM 3').font.size = Pt(8)
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Page ').font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
doc.core_properties.title = 'Estimated and Recorded Smartphone Screen Time Among BTech Students'
doc.core_properties.author = 'Prithviraj Gokule; Manish G S; Madhav Sunil'
doc.core_properties.subject = 'CES 2103 IA3 screen time analysis using simulated data'
doc.core_properties.comments = ''

def p(text='', size=None, align=None):
    para = doc.add_paragraph()
    if align is not None: para.alignment = align
    for i, part in enumerate(re.split(r'\*\*(.*?)\*\*', text)):
        run = para.add_run(part); run.bold = bool(i % 2)
        if size: run.font.size = Pt(size)
    return para

def heading(text, level=1):
    return doc.add_paragraph(text, style=f'Heading {level}')

def page(title):
    heading(title).paragraph_format.page_break_before = True

def caption(text):
    return doc.add_paragraph(text, style='Caption')

def table(headers, data, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for col, w in zip(t.columns, widths): col.width = Inches(w)
    for c, value in zip(t.rows[0].cells, headers): c.text = str(value)
    for values in data:
        for c, value in zip(t.add_row().cells, values): c.text = str(value)
    for i, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if i == 0: row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for j, (c, w) in enumerate(zip(row.cells, widths)):
            c.width = Inches(w); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = c._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                b = OxmlElement('w:'+edge)
                for k, v in [('val','single'),('sz','4'),('color','D9D9D9')]: b.set(qn('w:'+k),v)
                borders.append(b)
            pr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge, value in [('top',90),('bottom',90),('left',100),('right',100)]:
                m = OxmlElement('w:'+edge); m.set(qn('w:w'),str(value)); m.set(qn('w:type'),'dxa'); margins.append(m)
            pr.append(margins)
            if i == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'),'EAEAEA'); pr.append(shade)
            for para in c.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_before = para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.08
                for r in para.runs: r.font.size = Pt(10.5); r.bold = i == 0
    after = doc.add_paragraph(); after.paragraph_format.space_after = Pt(0); after.paragraph_format.line_spacing = .3
    return t

def figure(file, label, alt):
    para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True; para.paragraph_format.space_after = Pt(0)
    picture = para.add_run().add_picture(str(WORK / file), width=Inches(6.15))
    picture._inline.docPr.set('descr', alt)
    caption(label)

def code(text):
    lines = text.splitlines()
    for i, line in enumerate(lines):
        para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.left_indent = Inches(.15)
        para.paragraph_format.space_after = Pt(0 if i < len(lines)-1 else 9)
        para.paragraph_format.line_spacing = 1.05; para.paragraph_format.keep_with_next = i < len(lines)-1
        run = para.add_run(line); run.font.name = 'Courier New'; run.font.size = Pt(9)

def formula(text):
    para = p(text, align=WD_ALIGN_PARAGRAPH.CENTER)
    para.paragraph_format.space_after = Pt(8)

def link(url, size=10):
    para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_after = Pt(7)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), doc.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'),'205781'); pr.append(color)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'),str(size*2)); pr.append(sz)
    u = OxmlElement('w:u'); u.set(qn('w:val'),'single'); pr.append(u)
    r.append(pr); txt = OxmlElement('w:t'); txt.text = url; r.append(txt); h.append(r); para._p.append(h)
    return para

# Page 1: explain the question and provenance before presenting results.
p('**MANIPAL INSTITUTE OF TECHNOLOGY**',12,WD_ALIGN_PARAGRAPH.CENTER)
p('CES 2103 Data Analytics  |  IA3 Mini Project',11,WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph('Estimated and Recorded Smartphone Screen Time Among BTech Students',style='Title')
p('Team 3',11,WD_ALIGN_PARAGRAPH.CENTER)
table(['Team member','Registration number'],[
    ['Prithviraj Gokule','251090052818'],['Manish G S','251090052782'],['Madhav Sunil','251090051962']],[3.7,2.5])
heading('1 Introduction and Objective')
p('How accurately do B.Tech students estimate their daily smartphone use? We study this question by comparing an estimate with the average of seven daily phone records. Both measurements refer to **the same person, primary phone and seven completed days**. This keeps the comparison meaningful.')
p('Our objective is to measure the direction and size of estimation errors, then test whether the average difference is zero. Exploratory data analysis (EDA) describes the patterns. Confirmatory data analysis (CDA) tests the specific hypothesis using a paired t-test.')
heading('2 Dataset Source and Description')
p('The current source is **data/simulated/screen_time_simulated.csv**, created in Python with NumPy and the fixed random seed 20261017. It contains **120 simulated records**, of which **112 pass validation**. These are generated observations; no Google Forms responses were used for this analysis.')
p('Each row represents one simulated student with an estimate and seven daily totals. The generator includes different usage levels, day-to-day variation, a small weekend increase and estimates rounded to 15 minutes. It assumes mean underestimation of 20 minutes, so the resulting bias partly reflects the model design. Eight invalid records were added to exercise the cleaning rules.')
p('The planned survey uses the same fields: respondents estimate first, then enter daily totals for the same dates. Names, college identifiers and phone operating systems are not required. The current analysis demonstrates the method; conclusions about actual students require genuine responses.')

# Page 2: complete data dictionary and cleaning decisions.
page('3 Data Preprocessing')
p('We preserve the source CSV and create a separate cleaned table. **One row is one observation**: seven daily readings are averaged within that row rather than counted as seven independent students.')
heading('3.1 Variables and units',2)
caption('Table 1  Main variables used in the analysis')
table(['Variable','Meaning','Unit or type'],[
    ['Participant ID','Label assigned after import','Nominal'],
    ['Estimated daily time','Estimate before checking records','Minutes per day'],
    ['Seven daily totals','Usage for each completed day','Minutes per day'],
    ['Recorded daily average','Mean of all seven daily totals','Minutes per day'],
    ['Signed error','Recorded average minus estimate','Minutes per day'],
    ['Absolute error','Magnitude of the signed error','Minutes per day']],[1.65,3.0,1.55])
p('Durations arrive as H:MM text and are converted to minutes. For example, **4:30 = 270 minutes**, not 4.30 hours. The reference end date identifies the last day of the seven-day window. Five Yes/No fields record consent, B.Tech eligibility, first response, estimate-before-check confirmation and recording-scope confirmation.')
heading('3.2 Validation and exclusions',2)
p('A row is included only when all five confirmations are Yes, the date is valid, and the estimate plus every daily reading is present and between 0:00 and 24:00. Eight rows fail: one has a missing daily total, one an impossible duration, one an invalid date, and five each fail a different confirmation. Thus **120 − 8 = 112 complete pairs** remain.')
p('We do not replace missing durations with zero or an average. An incomplete week is excluded. Identical response patterns are flagged for review rather than automatically deleted, because matching values do not prove that the same person responded twice. The audit records each inclusion decision and its reason.')
heading('3.3 Calculated measures',2)
formula('Recorded daily average = sum of the seven daily totals / 7')
formula('Signed error = recorded daily average − estimated daily time')
p('**Positive error means underestimation; negative error means overestimation.** Absolute error is the size of this difference without its sign. For example, an estimate of 240 minutes and a recorded average of 300 minutes gives a signed error of +60 minutes and an absolute error of 60 minutes.')

# Page 3: one large figure, readable summary and interpretation.
page('4 Exploratory Data Analysis')
p(f'EDA describes the {n} included pairs before the hypothesis test. The mean gives the average, the median is the middle value, and the sample standard deviation (SD) describes variation between records. All durations below are in minutes per day.')
caption('Table 2  Descriptive statistics for the included records')
table(['Measure','Mean','Median','SD'],[
    [label, f2(summary[key]['mean']),f2(summary[key]['50%']),f2(summary[key]['std'])]
    for key,label in [('estimated_daily_minutes','Estimated time'),('recorded_daily_minutes','Recorded average'),('signed_error_minutes','Signed error'),('absolute_error_minutes','Absolute error')]
],[2.6,1.2,1.2,1.2])
heading('4.1 Comparing estimates with recorded averages',2)
figure('report_scatter.png','Figure 1  Each point represents one simulated record. The dashed line marks equal values.','Scatterplot of estimates against recorded averages with an equality line.')
p('Points above the dashed line show underestimation: the recorded average exceeds the estimate. Points below it show overestimation. Points close to the line have smaller errors. Although higher estimates generally accompany higher recorded use, the spread around the line shows that this association does not guarantee accurate estimates.')
p(f'The mean estimate is **{f2(summary["estimated_daily_minutes"]["mean"])} minutes**, while the recorded mean is **{f2(summary["recorded_daily_minutes"]["mean"])} minutes**. Their difference is **{f2(s["mean_difference"])} minutes per day**. The recorded median is {f2(summary["recorded_daily_minutes"]["50%"])} minutes, slightly below its mean, consistent with some higher-usage records pulling the mean upward.')
p('These summaries answer different questions. The two usage means describe overall levels, while the signed and absolute errors describe how closely each estimate matches its own recorded average.')

# Page 4: error distribution, individual accuracy and diagnostic evidence.
page('4 Exploratory Data Analysis Continued')
heading('4.2 Direction and size of the errors',2)
figure('report_errors.png','Figure 2  Errors occur on both sides of zero; the mean is positive.','Histogram of signed errors with zero and mean reference lines.')
p(f'**Mean signed error is {f2(s["mean_difference"])} minutes; mean absolute error is {f2(s["mean_absolute_error"])} minutes.** Overestimates and underestimates partly cancel when we calculate the mean signed error. For example, errors of +60 and −60 average to zero, although both estimates are wrong by 60 minutes.')
p(f'Using a descriptive 30-minute tolerance, {counts["Underestimated >30 min"]} records (47.3%) underestimate by more than 30 minutes, {counts["Overestimated >30 min"]} (19.6%) overestimate by more than 30 minutes, and {counts["Within 30 min"]} (33.0%) fall within 30 minutes. This tolerance is a project convention, not a health threshold or a validated accuracy standard.')
heading('4.3 Checking the differences before the test',2)
figure('report_qq.png','Figure 3  A normal Q–Q plot compares the observed differences with normal reference quantiles.','Normal Q-Q plot of signed errors, showing roughly linear central points and tail departures.')
p('The Q–Q plot is roughly linear in the centre, with departures in the tails. Three differences fall outside the 1.5 × IQR fences, where IQR is the spread of the middle half of the data. We retain these valid values and report the departures. The paired t-test concerns the distribution of differences; these checks do not prove normality or independence.')

# Page 5: inference, numerical working and responsible interpretation.
page('5 Confirmatory Data Analysis')
heading('5.1 Hypothesis and test selection',2)
p('We use a **two-sided paired t-test at α = 0.05**. Each estimate is matched to the same record’s seven-day average, so the analysis tests the within-record differences [1]. The two-sided choice allows either underestimation or overestimation.')
p('**H₀:** the population mean of recorded minus estimated time is zero. **H₁:** this mean difference is not zero. In this report, the population is the hypothetical one represented by the simulation.')
p('The test assumes independent participant differences and a sufficiently well-behaved distribution of differences. The previous page shows tail departures to consider when interpreting the result. A large sample cannot correct biased sampling or inaccurate entries.')
heading('5.2 Calculation and result',2)
p('The standard error (SE) describes uncertainty in the estimated mean. It differs from the SD, which describes variation across individual records. The test divides the mean difference by its SE.')
formula(f'SE = {s["sd_difference"]:.4f} / √{n} = {s["standard_error"]:.4f} minutes')
formula(f't = {s["mean_difference"]:.4f} / {s["standard_error"]:.4f} = {s["t_statistic"]:.4f}     df = {n} − 1 = {s["df"]}')
caption('Table 3  Paired t-test results for the simulated dataset')
table(['Result','Value'],[
    ['Complete pairs',n],['Mean difference',f'{s["mean_difference"]:.2f} minutes per day'],
    ['Test statistic and degrees of freedom',f't = {s["t_statistic"]:.3f}; df = {s["df"]}'],
    ['Two-sided p-value',f'{s["p_value"]:.5f}'],
    ['95% confidence interval',f'{s["ci_low"]:.2f} to {s["ci_high"]:.2f} minutes per day'],
    ['Standardised paired effect size',f'Cohen’s dz = {s["cohens_dz"]:.3f}']],[3.65,2.55])
heading('5.3 Interpretation',2)
p(f'Because **p = {s["p_value"]:.5f} < 0.05**, we reject zero mean difference for this simulated analysis. If H₀ and the test assumptions held, a t-statistic at least this far from zero would be unlikely. The p-value is not the probability that H₀ is true.')
p(f'The **95% confidence interval is {s["ci_low"]:.2f} to {s["ci_high"]:.2f} minutes per day** and excludes zero. It describes uncertainty in the population mean difference, not a range containing 95% of individual errors. Under repeated sampling with the model assumptions, 95% of intervals constructed this way would cover that mean.')
p(f'Cohen’s dz divides the mean difference by the SD of the differences: {s["mean_difference"]:.2f} / {s["sd_difference"]:.2f} = {s["cohens_dz"]:.3f}. The average bias is about 0.31 SD. The error in minutes remains the more direct measure of practical size.')

# Page 6: brief implementation extracts, source navigation and rerun contract.
page('6 Python Implementation and Reproducibility')
p('The workflow is **load, validate, calculate errors, explore, test and export**. The notebook explains each step; reusable Python functions perform the calculations.')
heading('6.1 Libraries and core calculations',2)
p('pandas organises the table; NumPy supports calculations. Matplotlib and Seaborn create plots, and SciPy provides the statistical test. CDA is the hypothesis-testing stage, rather than a separate Python library.')
p('Here, numeric_days contains the seven duration columns. axis=1 calculates across each row; skipna=False prevents an incomplete week from being silently averaged.')
code('frame["recorded_daily_minutes"] = frame[numeric_days].mean(\n    axis=1, skipna=False\n)\nframe["signed_error_minutes"] = (\n    frame["recorded_daily_minutes"]\n    - frame["estimated_daily_minutes"]\n)')
p('SciPy compares the matched columns [2]. Their order makes underestimation positive. nan_policy="raise" stops the test if missing values remain.')
code('test = stats.ttest_rel(\n    clean["recorded_daily_minutes"],\n    clean["estimated_daily_minutes"],\n    alternative="two-sided", nan_policy="raise"\n)')
p('The code checks that the formula-based t and p-value match SciPy, and exports the confidence interval, effect size and outlier flags. Automated checks cover time conversion, exclusions, source handling and statistical calculations.')
heading('6.2 Code files and saved outputs',2)
table(['File or folder','Role'],[
    ['run.py','Command-line entry point'],
    ['src/screentime/pipeline.py','Validation, EDA, paired test and exports'],
    ['src/screentime/simulate.py','Reproducible simulation generator'],
    ['notebooks/01_screen_time_analysis.ipynb','Explained analysis with outputs'],
    ['config/ and outputs/','Input settings and saved analysis results']],[3.15,3.05])
p('Run python run.py --config config/simulated.json. Outputs include cleaned data, an audit, descriptive statistics, figures and statistics.json. A manifest records the source checksum, settings and software versions.')
p('For a real survey, add data/raw/screen_time_real.csv, map form headings in config/real.json and rerun with that configuration. Refresh the notebook and report from the new outputs. Missing real data never trigger a fallback to simulation.')

# Page 7: synthesis, limitations and a visible clickable repository URL.
page('7 Findings and Conclusions')
heading('7.1 What the analysis shows',2)
p(f'The generated records show **average underestimation of {s["mean_difference"]:.2f} minutes per day**, but an **average absolute error of {s["mean_absolute_error"]:.2f} minutes**. This distinction is the main finding: a modest net bias can hide much larger errors in individual estimates.')
p('EDA shows errors in both directions and wide variation between records. CDA finds evidence of a nonzero mean difference within the simulated model. These results answer complementary questions: the paired test addresses average bias, while the plots and absolute errors describe the spread and size of the mismatches.')
heading('7.2 Limitations and next step',2)
p('The main limitation is the source. The data are simulated, and the generator already assumes some underestimation. The significant result therefore partly reflects those assumptions; it is not independent evidence about B.Tech students. The current project establishes a working analysis that can be applied to collected responses.')
p('A future survey shared through friends would be a convenience sample, so it would not represent all B.Tech students. Respondents might mistype device totals or change an estimate after checking. An anonymous form also cannot fully verify unique participation. Keeping the same phone, dates and instructions helps consistency but does not remove these limitations.')
p('The next step is to pilot the form, collect genuine responses and rerun the same analysis. The new results may support a different conclusion. Simulated and real records should remain separate, and every reported number and chart must be updated after a change of source.')
heading('7.3 Conclusion',2)
p('We implemented the full process from data validation to statistical interpretation. The analysis explains why mean bias alone is not enough to judge estimation accuracy: the direction, absolute size and spread of the errors all matter. With genuine survey data, this workflow can assess how closely respondents’ estimates match their phone records.')
heading('Project Repository',2)
link(REPO)
heading('References',2)
p('[1] NIST/SEMATECH. e-Handbook of Statistical Methods, Two-Sample t-Test for Equal Means, discussion of paired observations. Accessed 26 October 2026.',9.5,WD_ALIGN_PARAGRAPH.LEFT)
link('https://www.itl.nist.gov/div898/handbook/eda/section3/eda353.htm',9)
p('[2] SciPy documentation. scipy.stats.ttest_rel, related-sample t-test API. Accessed 26 October 2026.',9.5,WD_ALIGN_PARAGRAPH.LEFT)
link('https://docs.scipy.org/doc/scipy/reference/generated/scipy.stats.ttest_rel.html',9)

doc.save(OUT)
shutil.copyfile(OUT, ROOT / 'report/Screen_Time_IA3_Report.docx')
print(OUT)
